"""Markdown -> Word (.docx) 转换脚本
将 docs/农村自建房风险巡查管理功能需求规格说明书.md 转换为 docx
"""
import re
import sys
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


SRC = Path(r"e:\风险管控0618\docs\农村自建房风险巡查管理功能需求规格说明书.md")
DST = SRC.with_suffix(".docx")


def set_run_font(run, size=None, bold=None, color=None, name="微软雅黑"):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text_with_inline(paragraph, node):
    """递归处理 inline 元素，支持 strong/em/code 等"""
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size=11)
        return
    if not isinstance(node, Tag):
        return
    name = node.name.lower()
    if name in ('strong', 'b'):
        for child in node.children:
            if isinstance(child, NavigableString):
                run = paragraph.add_run(str(child))
                set_run_font(run, size=11, bold=True)
            else:
                # 递归后把最后添加的 run 加粗
                before = len(paragraph.runs)
                add_text_with_inline(paragraph, child)
                for r in paragraph.runs[before:]:
                    r.font.bold = True
        return
    if name in ('em', 'i'):
        for child in node.children:
            if isinstance(child, NavigableString):
                run = paragraph.add_run(str(child))
                set_run_font(run, size=11)
                run.font.italic = True
            else:
                before = len(paragraph.runs)
                add_text_with_inline(paragraph, child)
                for r in paragraph.runs[before:]:
                    r.font.italic = True
        return
    if name == 'code':
        text = node.get_text()
        run = paragraph.add_run(text)
        set_run_font(run, size=10, name='Consolas')
        run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        return
    if name == 'br':
        paragraph.add_run().add_break()
        return
    if name == 'a':
        text = node.get_text()
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color=(0x05, 0x63, 0xC1))
        run.font.underline = True
        return
    # 其他标签递归
    for child in node.children:
        add_text_with_inline(paragraph, child)


def add_heading(doc, level, text):
    p = doc.add_paragraph()
    sizes = {1: 22, 2: 18, 3: 15, 4: 13, 5: 12, 6: 11}
    size = sizes.get(level, 11)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(18)
    else:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    # 设置大纲级别，便于 Word 生成目录
    pPr = p._element.get_or_add_pPr()
    outlineLvl = OxmlElement('w:outlineLvl')
    outlineLvl.set(qn('w:val'), str(level - 1))
    pPr.append(outlineLvl)
    return p


def add_paragraph(doc, node, style=None, indent=0):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.75 * indent)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    for child in node.children:
        add_text_with_inline(p, child)
    return p


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '999999')
        tcBorders.append(b)
    tcPr.append(tcBorders)


def add_table(doc, table_tag):
    rows = table_tag.find_all('tr')
    if not rows:
        return
    ncols = 0
    for r in rows:
        cols = r.find_all(['th', 'td'], recursive=False)
        ncols = max(ncols, len(cols))
    if ncols == 0:
        return
    tbl = doc.add_table(rows=0, cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    for r_idx, row in enumerate(rows):
        cells = row.find_all(['th', 'td'], recursive=False)
        if not cells:
            continue
        tr = tbl.add_row()
        for c_idx in range(ncols):
            cell = tr.cells[c_idx]
            set_cell_border(cell)
            if c_idx < len(cells):
                tag = cells[c_idx]
                cell.text = ''
                para = cell.paragraphs[0]
                for child in tag.children:
                    add_text_with_inline(para, child)
                # 表头加粗
                if tag.name == 'th' or r_idx == 0:
                    for r in para.runs:
                        r.font.bold = True
                    # 表头底色
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'), 'D9E2F3')
                    cell._tc.get_or_add_tcPr().append(shd)
            else:
                cell.text = ''


def process_list(doc, list_tag, indent=0):
    ordered = list_tag.name.lower() == 'ol'
    for idx, li in enumerate(list_tag.find_all('li', recursive=False), start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75 * (indent + 1))
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.35
        marker = f"{idx}. " if ordered else "• "
        run = p.add_run(marker)
        set_run_font(run, size=11, bold=False)
        # li 下直接的文本和 inline
        for child in li.children:
            if isinstance(child, Tag) and child.name.lower() in ('ul', 'ol'):
                continue
            if isinstance(child, Tag) and child.name.lower() == 'p':
                for sub in child.children:
                    add_text_with_inline(p, sub)
            else:
                add_text_with_inline(p, child)
        # 嵌套列表
        for child in li.find_all(['ul', 'ol'], recursive=False):
            process_list(doc, child, indent + 1)


def process_block(doc, node):
    if isinstance(node, NavigableString):
        text = str(node).strip()
        if text:
            p = doc.add_paragraph()
            run = p.add_run(text)
            set_run_font(run, size=11)
        return
    if not isinstance(node, Tag):
        return
    name = node.name.lower()
    if name in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
        add_heading(doc, int(name[1]), node.get_text().strip())
    elif name == 'p':
        add_paragraph(doc, node)
    elif name == 'table':
        add_table(doc, node)
    elif name in ('ul', 'ol'):
        process_list(doc, node)
    elif name == 'blockquote':
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(node.get_text().strip())
        set_run_font(run, size=11)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    elif name == 'pre':
        code_text = node.get_text()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(code_text)
        set_run_font(run, size=10, name='Consolas')
    elif name == 'hr':
        p = doc.add_paragraph()
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:color'), '999999')
        pBdr.append(bottom)
        pPr.append(pBdr)
    else:
        # 其他容器递归
        for child in node.children:
            process_block(doc, child)


def main():
    if not SRC.exists():
        print(f"源文件不存在: {SRC}", file=sys.stderr)
        sys.exit(1)
    md_text = SRC.read_text(encoding='utf-8')
    html = markdown.markdown(
        md_text,
        extensions=['extra', 'tables', 'fenced_code', 'toc', 'sane_lists']
    )
    soup = BeautifulSoup(html, 'html.parser')

    doc = Document()
    # 默认正文样式
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(11)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 页面设置 A4
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    body = soup.body if soup.body else soup
    for child in body.children:
        process_block(doc, child)

    doc.save(DST)
    print(f"转换完成: {DST}")


if __name__ == '__main__':
    main()
