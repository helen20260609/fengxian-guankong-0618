# -*- coding: utf-8 -*-
"""
房屋安全管理——房屋建筑档案 详情页字段表生成器
依据 pages/house-arch-detail.html 的实际渲染逻辑提取所有字段
输出: docs/房屋建筑档案字段表.xlsx
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

HEADER_FILL = PatternFill("solid", fgColor="1a73e8")
MODULE_FILL = PatternFill("solid", fgColor="e8f0fe")
SUBTAB_FILL = PatternFill("solid", fgColor="f6f9ff")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
MODULE_FONT = Font(name="微软雅黑", size=11, bold=True, color="1a73e8")
SUBTAB_FONT = Font(name="微软雅黑", size=10, bold=True, color="1557b0")
BODY_FONT = Font(name="微软雅黑", size=10)
BORDER = Border(
    left=Side(style="thin", color="dadce0"),
    right=Side(style="thin", color="dadce0"),
    top=Side(style="thin", color="dadce0"),
    bottom=Side(style="thin", color="dadce0"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "房屋建筑档案字段表"

headers = ["一级模块", "二级Tab", "序号", "字段名称", "字段键/来源", "数据类型", "示例值", "备注"]
ws.append(headers)
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = CENTER
    cell.border = BORDER

FIELDS = [
    # ======== 顶部状态栏 ========
    ("顶部状态栏", "(全局)", [
        ("房屋标题", "name + owner", "字符串", "陈连明自住民居（李志强）", "页面顶栏"),
        ("房屋编号", "no", "字符串", "NF-2025-00123", "页面顶栏右侧"),
        ("风险等级", "riskLevel", "字符串", "第一类/第二类/第三类", "对应 safe/warning/danger"),
        ("治理状态", "governStatus", "字符串", "待整治/整治中/已整治/逾期未整治", ""),
        ("当前措施", "currentMeasure", "字符串", "管理措施（停止使用） + 工程措施（维修加固）", ""),
        ("责任人", "managerName + managerPhone", "字符串", "李志强 138****1234", ""),
    ]),
    # ======== 1. 基础信息 basic ========
    ("基础信息", "基本信息 jbxx", [
        ("房屋编号", "no", "字符串", "NF-2025-00123", "主键"),
        ("房屋名称", "name", "字符串", "陈连明自住民居", ""),
        ("房屋性质", "houseNature", "字符串", "自住民居", ""),
        ("所在区域", "region", "字符串", "奉贤区西渡街道", ""),
        ("路号信息", "roadNo", "字符串", "西闸公路二一三九弄35号", ""),
        ("不需要排查原因", "noSurveyReason", "字符串", "—", ""),
        ("产权人 / 使用人信息", "owner", "字符串", "陈连明", ""),
        ("身份证号码", "idCard", "字符串", "31022619680714051X", ""),
        ("房屋用途", "usageType", "字符串", "自住", ""),
        ("是否同时自住", "isSelfLive", "布尔", "是", ""),
        ("具体用途", "specificUsage", "字符串", "日常居住", ""),
        ("人员聚集场所周边", "crowdAround", "布尔", "否", ""),
        ("其他人员密集场所周边", "otherCrowdAround", "布尔", "否", ""),
        ("土地性质", "landNature", "字符串", "集体土地", ""),
        ("集体土地", "collectiveLand", "字符串", "宅基地", ""),
        ("所在区域", "region", "字符串", "奉贤区西渡街道", "重复字段"),
        ("建筑层数", "floors", "字符串", "3层", ""),
        ("建筑面积 (m²)", "buildingArea", "字符串", "237.3", ""),
        ("建成时间", "builtTime", "字符串", "2001-2010年", "区间"),
        ("具体年份", "builtYear", "字符串", "2007年", ""),
    ]),
    ("基础信息", "宅基地信息 zjd", [
        ("宅基地面积", "homesteadArea", "字符串", "132㎡", ""),
        ("审批情况", "homesteadApproval", "字符串", "自然资源部门审批通过", ""),
        ("权属证明", "homesteadCertNo", "字符串", "ZD-20070419", ""),
    ]),
    ("基础信息", "建设情况 jsqk", [
        ("常住或使用人数超过 10 人", "over10People", "字符串", "否", ""),
        ("设计方式", "designMode", "字符串", "有专业设计", ""),
        ("建造方式", "buildMode", "字符串", "有资质施工队伍", ""),
        ("结构类型", "structureType", "字符串", "砌体结构", ""),
        ("楼屋盖", "roofType", "字符串", "现浇板", ""),
        ("是否改扩建", "expansionStatus", "字符串", "否", ""),
        ("是否装修", "decorationStatus", "字符串", "否", ""),
    ]),
    ("基础信息", "管理情况 glqk", [
        ("已取得的行政许可", "permit", "字符串", "宅基地批准书", ""),
        ("违法建设和违法审批", "illegalBuild", "字符串", "否", ""),
    ]),
    # ======== 2. 排查信息 survey ========
    ("排查信息", "排查情况 survey1", [
        ("结构状况", "survey.structCondition", "字符串", "C级（局部危险）", ""),
        ("损伤具体部位", "survey.damagePart", "字符串", "东侧承重墙", ""),
        ("增加堆载（二层及以上堆放大量重物）", "survey.stackLoad", "字符串", "否", ""),
        ("其他使用安全隐患", "survey.otherHazard", "字符串", "暂无", ""),
        ("初步判定", "survey.preliminaryJudgment", "字符串", "立即停止使用", ""),
        ("位置标注", "survey.locationMark", "字符串", "—", ""),
        ("房屋安全责任人", "survey.safetyOwner", "字符串", "陈连明", ""),
        ("联系方式", "survey.safetyOwnerPhone", "字符串", "138****1234", ""),
        ("安全监管责任人", "survey.safetySupervisor", "字符串", "王建军", ""),
        ("联系方式", "survey.safetySupervisorPhone", "字符串", "139****5678", ""),
        ("排查照片", "survey.photos", "数组", "[照片]", ""),
        ("安全鉴定：是否整治前鉴定", "survey.needAppraisal", "字符串", "否", ""),
        ("未进行鉴定房屋原因", "survey.noAppraisalReason", "字符串", "资金尚未到位", ""),
        ("上传证明文件", "survey.proofFiles", "数组", "[文件]", ""),
        ("排查备注", "survey.remark", "字符串", "东侧承重墙纵向裂缝长约1.2m", ""),
        ("更新时间", "survey.updateTime", "日期时间", "2024-06-12 09:30", ""),
    ]),
    ("排查信息", "建房现场安全自查情况 survey2", [
        ("建房类型", "selfCheck.buildType", "字符串", "-", ""),
        ("实施路径", "selfCheck.implementPath", "字符串", "-", ""),
        ("目前建房进度", "selfCheck.buildProgress", "字符串", "-", ""),
        ("是否属于初判存在隐患的农村自建房的解危工程措施", "selfCheck.isReliefMeasure", "字符串", "-", ""),
        ("自查现场是否存在施工安全隐患", "selfCheck.hasHazard", "字符串", "-", ""),
        ("隐患是否已消除", "selfCheck.hazardCleared", "字符串", "-", ""),
        ("消除隐患的具体措施", "selfCheck.clearMeasure", "字符串", "-", ""),
        ("备注", "selfCheck.remark", "字符串", "-", ""),
        ("建房现场照片", "selfCheck.photos", "数组", "[照片]", ""),
        ("填报人", "selfCheck.reporter", "字符串", "-", ""),
        ("填报时间", "selfCheck.reportTime", "日期时间", "-", ""),
    ]),
    ("排查信息", "房屋用途现场情况 survey3", [
        ("房屋用途", "usageSurvey.usageType", "字符串", "-", ""),
        ("当前使用情况", "usageSurvey.currentUsage", "字符串", "-", ""),
        ("当前房屋用途", "usageSurvey.currentHouseUsage", "字符串", "-", ""),
        ("填报人", "usageSurvey.reporter", "字符串", "-", ""),
        ("填报时间", "usageSurvey.reportTime", "日期时间", "-", ""),
    ]),
    # ======== 3. 安全鉴定 assessment ========
    ("安全鉴定", "(单Tab)", [
        ("隐患动态整治状态", "assessment.rectifyStatus", "字符串", "整治中", ""),
        ("鉴定机构名称", "assessment.appraisalOrg", "字符串", "奉贤区房屋安全鉴定中心", ""),
        ("统一社会信用代码", "assessment.orgCode", "字符串", "91310120MA1K123456", ""),
        ("鉴定时间", "assessment.appraisalTime", "日期", "2024-06-20", ""),
        ("鉴定负责人", "assessment.appraiser", "字符串", "张伟", ""),
        ("鉴定结论", "assessment.conclusion", "字符串", "C级（局部危险）", "A/B/C/D"),
        ("阶段性标识", "assessment.stageMark", "字符串", "阶段性鉴定", ""),
        ("鉴定报告", "assessment.report", "字符串", "—", "附件"),
        ("备注", "assessment.remark", "字符串", "东侧承重墙需加固", ""),
        ("填报人", "assessment.reporter", "字符串", "李志强", ""),
        ("填报时间", "assessment.reportTime", "日期时间", "2024-06-21 10:00", ""),
    ]),
    # ======== 4. 整治计划 plan ========
    ("整治计划", "(单Tab)", [
        ("计划开始时间", "plan.startTime", "日期", "2024-07-05", ""),
        ("计划完成时间", "plan.endTime", "日期", "2024-09-30", ""),
        ("负责人信息", "plan.managerName", "字符串", "李志强", ""),
        ("联系电话", "plan.managerPhone", "字符串", "138-1234-5678", ""),
        ("整治方案", "plan.rectifyPlan", "字符串", "承重墙裂缝灌浆加固、屋面防水重做", ""),
        ("是否采取硬隔离", "plan.hardIsolation", "布尔", "否", ""),
        ("未硬隔离原因说明", "plan.noHardIsolationReason", "字符串", "资金尚未到位，暂未实施硬隔离", ""),
        ("长期管控照片", "plan.longTermPhotos", "数组", "[照片]", ""),
        ("告知承诺书", "plan.noticeLetter", "字符串", "—", ""),
        ("书面计划", "plan.writtenPlan", "字符串", "NF-2024-00001-FA", "档案编号"),
        ("备注", "plan.remark", "字符串", "计划已审核，待资金到位后启动施工", ""),
        ("填报人", "plan.reporter", "字符串", "李志强", ""),
        ("填报时间", "plan.reportTime", "日期时间", "2024-07-05 14:30", ""),
    ]),
    # ======== 5. 整治措施 measure ========
    ("整治措施", "管理措施 manage", [
        ("措施类型", "measure.manage.type", "字符串", "停止使用", "停止使用/停止经营/封控警示/人员撤离/持续监控"),
        ("实施部位", "measure.manage.position", "字符串", "整栋房屋", ""),
        ("开始时间", "measure.manage.startTime", "日期", "2024-07-05", ""),
        ("计划结束时间", "measure.manage.endTime", "日期", "2024-09-30", ""),
        ("实施要求", "measure.manage.requirement", "字符串", "立即停止使用危险区域，设置围挡和警示标识", ""),
        ("责任单位", "measure.manage.org", "字符串", "街道安全管理办公室", ""),
        ("责任人", "measure.manage.manager", "字符串", "李志强", ""),
        ("联系电话", "measure.manage.phone", "字符串", "138-1234-5678", ""),
        ("管理措施", "measure.manage.measure", "字符串", "停止使用、设置警示标识", ""),
        ("管理措施实施照片", "measure.manage.implementPhotos", "数组", "[照片]", ""),
        ("管理措施完成照片（最少上传两张照片）", "measure.manage.completePhotos", "数组", "[照片]", "最少2张"),
        ("措施是否完成", "measure.manage.isComplete", "布尔", "是", ""),
        ("措施效果评估", "measure.manage.effectEval", "字符串", "风险已有效控制，无新增变形", ""),
        ("风险是否已有效控制", "measure.manage.isControlled", "布尔", "是", ""),
        ("变更或延长申请审批", "measure.manage.changeApproval", "布尔", "是", ""),
        ("备注", "measure.manage.remark", "字符串", "已设置围挡和警示标识", ""),
        ("填报人", "measure.manage.reporter", "字符串", "李志强", ""),
        ("填报时间", "measure.manage.reportTime", "日期时间", "2024-07-05 14:30", ""),
        ("停止使用范围", "measure.manage.extra.stopUseScope", "字符串", "整栋房屋", "措施类型=停止使用时"),
        ("预计恢复使用时间", "measure.manage.extra.recoverTime", "日期", "2024-09-30", "措施类型=停止使用时"),
        ("经营项目", "measure.manage.extra.bizItem", "字符串", "—", "措施类型=停止经营时"),
        ("停止经营时间", "measure.manage.extra.stopBizTime", "日期", "—", "措施类型=停止经营时"),
        ("警戒范围", "measure.manage.extra.alertScope", "字符串", "—", "措施类型=封控警示时"),
        ("警示标识数量", "measure.manage.extra.alertSignCount", "数字", "—", "措施类型=封控警示时"),
        ("巡查频次", "measure.manage.extra.patrolFreq", "字符串", "—", "措施类型=封控警示/持续监控时"),
        ("撤离人数", "measure.manage.extra.evacCount", "数字", "—", "措施类型=人员撤离时"),
        ("安置方式", "measure.manage.extra.evacPlace", "字符串", "—", "措施类型=人员撤离时"),
        ("撤离确认人", "measure.manage.extra.evacConfirmer", "字符串", "—", "措施类型=人员撤离时"),
        ("撤离时间", "measure.manage.extra.evacTime", "日期", "—", "措施类型=人员撤离时"),
        ("监控设备", "measure.manage.extra.monitorDevice", "字符串", "—", "措施类型=持续监控时"),
        ("监控频次", "measure.manage.extra.monitorFreq", "字符串", "—", "措施类型=持续监控时"),
        ("巡查责任人", "measure.manage.extra.patrolOwner", "字符串", "—", "措施类型=持续监控时"),
    ]),
    ("整治措施", "工程措施 project", [
        ("工程措施", "measure.project.projectMeasure", "字符串", "维修加固", ""),
        ("措施类型", "measure.project.type", "字符串", "承重墙裂缝灌浆加固、屋面防水重做", ""),
        ("实施范围", "measure.project.scope", "字符串", "整栋房屋主体及屋面", ""),
        ("实施拟采用的政策路径", "measure.project.policyPath", "字符串", "区级修缮工程补贴", ""),
        ("其他实施拟采用政策路径", "measure.project.otherPolicy", "字符串", "—", ""),
        ("细化工程措施方案", "measure.project.detailPlan", "字符串", "采用压力灌浆封闭裂缝，屋面重新铺设防水卷材", ""),
        ("开工时间", "measure.project.startTime", "日期", "2024-08-01", ""),
        ("竣工时间", "measure.project.endTime", "日期", "2024-09-30", ""),
        ("施工单位", "measure.project.constructionOrg", "字符串", "奉贤区房屋修缮工程公司", ""),
        ("施工负责人", "measure.project.constructionManager", "字符串", "王建军", ""),
        ("联系电话", "measure.project.phone", "字符串", "139-5678-1234", ""),
        ("工程概算（万元）", "measure.project.budget", "数字", "12.5", ""),
        ("资金来源", "measure.project.fundSource", "字符串", "区级修缮补贴60%+业主自筹40%", ""),
        ("经审核的施工方案文件", "measure.project.planFiles", "数组", "[文件]", ""),
        ("设计图纸", "measure.project.designFiles", "数组", "[文件]", ""),
        ("施工组织设计", "measure.project.orgDesignFiles", "数组", "[文件]", ""),
        ("安全专项方案", "measure.project.safetyFiles", "数组", "[文件]", ""),
        ("基础验收记录", "measure.project.baseAcceptFiles", "数组", "[文件]", ""),
        ("主体验收记录", "measure.project.mainAcceptFiles", "数组", "[文件]", ""),
        ("材料进场验收记录", "measure.project.materialFiles", "数组", "[文件]", ""),
        ("工程措施实施照片", "measure.project.implementPhotos", "数组", "[照片]", ""),
        ("竣工验收报告", "measure.project.completeReportFiles", "数组", "[文件]", ""),
        ("验收结论", "measure.project.acceptConclusion", "字符串", "—", ""),
        ("验收人员", "measure.project.acceptPerson", "字符串", "—", ""),
        ("验收日期", "measure.project.acceptDate", "日期", "—", ""),
        ("整改后照片", "measure.project.rectifyPhotos", "数组", "[照片]", ""),
        ("工程措施完成照片", "measure.project.completePhotos", "数组", "[照片]", ""),
        ("措施是否完成", "measure.project.isComplete", "布尔", "否", ""),
        ("备注", "measure.project.remark", "字符串", "已进场施工，进度约20%", ""),
        ("填报人", "measure.project.reporter", "字符串", "王建军", ""),
        ("填报时间", "measure.project.reportTime", "日期时间", "2024-08-01 08:00", ""),
    ]),
    # ======== 6. 隐患控制情况 control ========
    ("隐患控制情况", "(单Tab)", [
        ("隐患动态整治状态", "control.rectifyStatus", "字符串", "整治中", ""),
        ("安全隐患是否得到控制", "control.isControlled", "字符串", "是", ""),
        ("备注", "control.remark", "字符串", "房屋已停用，设置围挡警示", ""),
        ("填报人", "control.reporter", "字符串", "王建军", ""),
        ("填报时间", "control.reportTime", "日期时间", "2024-08-05 09:00", ""),
    ]),
    # ======== 7. 阶段性验收 accept ========
    ("阶段性验收", "(单Tab)", [
        ("验收方式", "accept.acceptMethod", "字符串", "第三方专业机构验收", ""),
        ("整治完成后是否进行了安全鉴定", "accept.hasAppraisal", "字符串", "是", ""),
        ("经专业第三方复核判定", "accept.thirdPartyResult", "字符串", "B级（安全）", ""),
        ("姓名", "accept.contactName", "字符串", "张伟", ""),
        ("联系电话", "accept.contactPhone", "字符串", "136-0000-1234", ""),
        ("整治后照片", "accept.afterPhotos", "数组", "[照片]", ""),
        ("相关资料", "accept.relatedFiles", "数组", "[文件]", ""),
        ("备注", "accept.remark", "字符串", "待整治完成后提交验收", ""),
        ("填报人", "accept.reporter", "字符串", "李志强", ""),
        ("填报时间", "accept.reportTime", "日期时间", "2024-09-30 10:00", ""),
        ("提交人", "accept.auditList[].submitter", "字符串", "李志强", "审核记录"),
        ("提交时间", "accept.auditList[].submitTime", "日期时间", "2024-09-28 09:00", "审核记录"),
        ("审核时间", "accept.auditList[].auditTime", "日期时间", "2024-09-30 10:00", "审核记录"),
        ("审核人", "accept.auditList[].auditor", "字符串", "王建国", "审核记录"),
        ("审核结论", "accept.auditList[].conclusion", "字符串", "通过/待审核", "审核记录"),
        ("审核报告", "accept.auditList[].report", "字符串", "查看报告", "审核记录"),
        ("备注", "accept.auditList[].remark", "字符串", "符合验收标准", "审核记录"),
    ]),
    # ======== 8. 验收销项 close ========
    ("验收销项", "(单Tab)", [
        ("隐患动态整治状态", "close.rectifyStatus", "字符串", "-", ""),
        ("工程措施", "close.projectMeasure", "字符串", "-", ""),
        ("工程措施实施照片", "close.implementPhotos", "数组", "[照片]", ""),
        ("工程措施完成照片", "close.completePhotos", "数组", "[照片]", ""),
        ("完工证明", "close.completeCert", "字符串", "-", ""),
        ("专业意见", "close.professionalOpinion", "字符串", "-", ""),
        ("专业认定意见材料", "close.professionalFiles", "数组", "[文件]", ""),
        ("销项方式", "close.closeMethod", "字符串", "-", ""),
        ("已完成整治且安全隐患已消除", "close.isCleared", "字符串", "-", ""),
        ("整治完成后是否进行了安全鉴定", "close.hasAppraisal", "字符串", "-", ""),
        ("经专业第三方复核判定", "close.thirdPartyResult", "字符串", "-", ""),
        ("姓名", "close.contactName", "字符串", "-", ""),
        ("联系电话", "close.contactPhone", "字符串", "-", ""),
        ("整治后照片", "close.afterPhotos", "数组", "[照片]", ""),
        ("相关资料", "close.relatedFiles", "数组", "[文件]", ""),
        ("是否更新房屋信息", "close.updateHouseInfo", "布尔", "-", ""),
        ("备注", "close.remark", "字符串", "-", ""),
        ("填报人", "close.reporter", "字符串", "-", ""),
        ("填报时间", "close.reportTime", "日期时间", "-", ""),
        ("提交人", "close.auditList[].submitter", "字符串", "张工", "审核记录"),
        ("提交时间", "close.auditList[].submitTime", "日期时间", "2024-10-05 09:30", "审核记录"),
        ("审核时间", "close.auditList[].auditTime", "日期时间", "2024-10-07 14:00", "审核记录"),
        ("审核人", "close.auditList[].auditor", "字符串", "李建国", "审核记录"),
        ("审核结论", "close.auditList[].conclusion", "字符串", "通过/待审核", "审核记录"),
        ("审核报告", "close.auditList[].report", "字符串", "查看报告", "审核记录"),
        ("备注", "close.auditList[].remark", "字符串", "销项资料齐全", "审核记录"),
    ]),
    # ======== 9. 隐患动态记录 record ========
    ("隐患动态记录", "(单Tab)", [
        ("隐患判定情况", "record.hazardStatus", "字符串", "-", "可多条"),
        ("新增隐患原因", "record.newHazardReason", "字符串", "-", ""),
        ("填报人", "record.reporter", "字符串", "-", ""),
        ("填报时间", "record.reportTime", "日期时间", "-", ""),
    ]),
    # ======== 10. 整治档案 archive ========
    ("整治档案", "(单Tab)", [
        ("档案类型", "archive.archiveType", "字符串", "整治方案/实施记录/验收材料", "可多条"),
        ("档案编号", "archive.archiveNo", "字符串", "NF-2024-00001-FA", ""),
        ("归档时间", "archive.archiveTime", "日期", "2024-07-08", ""),
        ("归档状态", "archive.archiveStatus", "字符串", "已归档/待归档", ""),
    ]),
]

current_row = 2
for module, subtab, fields in FIELDS:
    module_start = current_row
    for idx, (name, key, dtype, example, remark) in enumerate(fields, 1):
        ws.cell(row=current_row, column=1, value=module)
        ws.cell(row=current_row, column=2, value=subtab)
        ws.cell(row=current_row, column=3, value=idx)
        ws.cell(row=current_row, column=4, value=name)
        ws.cell(row=current_row, column=5, value=key)
        ws.cell(row=current_row, column=6, value=dtype)
        ws.cell(row=current_row, column=7, value=example)
        ws.cell(row=current_row, column=8, value=remark)
        current_row += 1
    module_end = current_row - 1

    if module_end > module_start:
        ws.merge_cells(start_row=module_start, start_column=1, end_row=module_end, end_column=1)
        ws.merge_cells(start_row=module_start, start_column=2, end_row=module_end, end_column=2)

    for r in range(module_start, module_end + 1):
        ws.cell(row=r, column=1).border = BORDER
        ws.cell(row=r, column=1).fill = MODULE_FILL
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=2).fill = SUBTAB_FILL

    mc = ws.cell(row=module_start, column=1)
    mc.font = MODULE_FONT
    mc.alignment = CENTER
    sc = ws.cell(row=module_start, column=2)
    sc.font = SUBTAB_FONT
    sc.alignment = CENTER

for row in ws.iter_rows(min_row=2, max_row=current_row - 1):
    for cell in row:
        cell.border = BORDER
        if cell.column in (1, 2, 3, 6):
            cell.alignment = CENTER
        else:
            cell.alignment = LEFT
        if cell.column >= 3:
            cell.font = BODY_FONT

col_widths = [16, 22, 6, 32, 40, 10, 36, 28]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"

out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "房屋建筑档案字段表.xlsx")
wb.save(out_path)
print(f"已生成: {out_path}")
print(f"共 {current_row - 2} 个字段, {len(FIELDS)} 个模块/Tab")
