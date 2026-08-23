# -*- coding: utf-8 -*-
"""
将 docs/农村自建房风险档案总览功能需求规格说明书_20260823.md 转换为规范 Word 文档
- 封面页
- 自动目录（TOC 域，打开 Word 后按 Ctrl+A 再 F9 更新生成）
- 一级/二级/三级标题样式
- 表格带边框、表头底纹
- 页眉、页脚页码
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE_DIR, "docs", "农村自建房风险档案总览功能需求规格说明书_20260823.md")
DST = os.path.splitext(SRC)[0] + ".docx"


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), 'BFBFBF')
        tcBorders.append(b)
    tcPr.append(tcBorders)


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def add_toc(doc):
    para = doc.add_paragraph()
    run = para.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    hint = OxmlElement('w:t')
    hint.text = "（目录将在 Word 中打开后自动生成；若未生成，请全选后按 F9）"
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(hint)
    run._r.append(fld_end)


def add_page_number_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_header_text(doc, text):
    section = doc.sections[0]
    header = section.header
    para = header.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = para.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')


def style_setup(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    for level, (size, color) in {
        1: (18, RGBColor(0x1A, 0x73, 0xE8)),
        2: (15, RGBColor(0x15, 0x57, 0xB0)),
        3: (13, RGBColor(0x33, 0x33, 0x33)),
    }.items():
        st = doc.styles[f'Heading {level}']
        st.font.name = 'Times New Roman'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')


def add_cover(doc, title, subtitle, date_str):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.name = 'Times New Roman'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.size = Pt(16)
    r2.font.name = 'Times New Roman'
    r2._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    for _ in range(10):
        doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(date_str)
    r3.font.size = Pt(14)
    r3.font.name = 'Times New Roman'
    r3._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    doc.add_page_break()


def parse_inline(paragraph, text):
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            paragraph.add_run(part)


def add_table(doc, header_cells, rows):
    ncols = len(header_cells)
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for i, txt in enumerate(header_cells):
        hdr[i].text = ''
        para = hdr[i].paragraphs[0]
        run = para.add_run(txt)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = 'Times New Roman'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        shade_cell(hdr[i], '1A73E8')
        set_cell_border(hdr[i])
    for row in rows:
        cells = table.add_row().cells
        for i in range(ncols):
            txt = row[i] if i < len(row) else ''
            cells[i].text = ''
            para = cells[i].paragraphs[0]
            run = para.add_run(txt)
            run.font.size = Pt(10.5)
            run.font.name = 'Times New Roman'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            set_cell_border(cells[i])
    doc.add_paragraph()


def convert(md_path, out_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    style_setup(doc)
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    base = os.path.basename(md_path)
    title = re.sub(r'_\d{8}\.md$', '', base)
    add_cover(doc, title, '功能需求规格说明书', '2026年8月23日')

    p = doc.add_paragraph('目  录')
    p.style = doc.styles['Heading 1']
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_toc(doc)
    doc.add_page_break()

    add_header_text(doc, title)
    add_page_number_footer(doc)

    i = 0
    in_code = False
    code_buf = []
    table_header = None
    table_rows = []
    table_sep_seen = False

    def flush_table():
        nonlocal table_header, table_rows, table_sep_seen
        if table_header is not None:
            add_table(doc, table_header, table_rows)
        table_header = None
        table_rows = []
        table_sep_seen = False

    def flush_code():
        nonlocal code_buf
        if code_buf:
            para = doc.add_paragraph()
            run = para.add_run('\n'.join(code_buf))
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
            code_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code:
                in_code = False
                flush_code()
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            is_sep = all(re.fullmatch(r':?-{3,}:?', c or '---') for c in cells)
            if table_header is None:
                table_header = cells
                table_sep_seen = False
            elif not table_sep_seen and is_sep:
                table_sep_seen = True
            else:
                table_rows.append(cells)
            i += 1
            continue
        else:
            flush_table()

        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and text == title:
                i += 1
                continue
            hlevel = min(level, 3)
            p = doc.add_heading('', level=hlevel)
            parse_inline(p, text)
            i += 1
            continue

        if re.fullmatch(r'-{3,}|\*{3,}|_{3,}', stripped):
            i += 1
            continue

        m = re.match(r'^[-*+]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            parse_inline(p, m.group(1))
            i += 1
            continue

        m = re.match(r'^\d+[.)]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            parse_inline(p, m.group(1))
            i += 1
            continue

        m = re.match(r'^>\s?(.*)$', stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            run = p.add_run(m.group(1))
            run.italic = True
            run.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        parse_inline(p, stripped)
        i += 1

    flush_table()
    flush_code()
    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    out = convert(SRC, DST)
    print(f'已生成: {out}')
